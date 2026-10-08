import openpyxl, json, datetime, re, sys
src=sys.argv[1]
ws=openpyxl.load_workbook(src,data_only=True).worksheets[0]
hdr={ws.cell(15,c).value:c for c in range(1,34) if ws.cell(15,c).value}
changes=[]; projects=[]; variations=[]
def d(v): return v.date().isoformat() if isinstance(v,datetime.datetime) else None
def s(v): return str(v).strip() if v is not None else None
for r in range(16, ws.max_row+1):
    name=ws.cell(r,hdr['PROJECT NAME']).value
    if not name: continue
    row={}
    def take(col,key,kind='text'):
        raw=ws.cell(r,hdr[col]).value
        if kind=='date':
            val=d(raw)
            if raw is not None and val is None:
                changes.append((r,col,repr(raw),'blank','not a date')); 
        else:
            val=s(raw)
            if val in ('_','-',''): 
                if raw not in (None,''): changes.append((r,col,repr(raw),'blank','placeholder'))
                val=None
            elif raw is not None and str(raw)!=val: changes.append((r,col,repr(raw),repr(val),'trimmed whitespace/newline'))
        row[key]=val
    take('ENQ','enquiry_no'); take('PROJECT NAME','name'); take('CLIENT','client')
    take('MAIN CONTRACTOR','contractor'); take('PROJECT MANAGER','pm'); take('PROJECT DESIGNER','designer')
    take('TECHNICAL DESIGNER','technical_designer'); take('STATUS','status')
    take('RECEIVED INQUIRY','received_date','date'); take('PROJECT START','start_date','date')
    take('TARGET COMPLETION','target_date','date'); take('DATE COMPLETED','completed_date','date')
    row['name']=re.sub(r'\s+',' ',row['name'])
    if row['pm'] and '/' in row['pm']:
        a,b=[x.strip() for x in row['pm'].split('/')]
        changes.append((r,'PROJECT MANAGER',row['pm'],f'{a} (+ second PM {b})','shared PM')); row['pm']=a; row['pm_second']=b
    else: row['pm_second']=None
    if row['contractor']=='115': changes.append((r,'MAIN CONTRACTOR','115','115 — CONFIRM','looks like a typo'))
    if row['status']=='COMPLETED' and not row['completed_date']:
        changes.append((r,'DATE COMPLETED','blank','blank','completed but no completion date — Aftab to fill'))
    sq=ws.cell(r,hdr['SIGNED QUOTE']).value
    row['signed_quote']=float(sq) if isinstance(sq,(int,float)) else 0.0
    for i in range(1,16):
        v=ws.cell(r,hdr[f'VO{i}']).value
        if isinstance(v,(int,float)) and v!=0:
            variations.append({'enquiry_no':row['enquiry_no'],'project_name':row['name'],'vo_no':i,'value':float(v)})
    row['excel_row']=r
    projects.append(row)
json.dump({'projects':projects,'variations':variations},open('import_data.json','w'),indent=1,default=str)
with open('Import_Report_v1.md','w') as f:
    f.write(f'# Import report — CONCEPT_5-_PROJECT_2026.xlsx\n\n{len(projects)} projects, {len(variations)} variation orders, pipeline value AED {sum(p["signed_quote"] for p in projects)+sum(v["value"] for v in variations):,.2f}\n\n')
    f.write('| Excel row | Project | Column | Was | Now | Why |\n|---|---|---|---|---|---|\n')
    for r,col,a,b,why in changes:
        nm=next(p['name'] for p in projects if p['excel_row']==r)
        f.write(f'| {r} | {nm[:35]} | {col} | {a} | {b} | {why} |\n')
print(len(projects),'projects',len(variations),'VOs',len(changes),'changes')
