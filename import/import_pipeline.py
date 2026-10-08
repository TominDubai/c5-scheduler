"""Turn Ramus's Estimation-Tracker.xlsx (Pipeline sheet) into SQL for scheduler.projects.

usage: python3 -I import_pipeline.py Estimation-Tracker.xlsx > pipeline_import.sql
Writes pipeline_report.md alongside.
"""
import sys, json, collections, datetime as dt
import openpyxl

SRC = sys.argv[1]
CONTRACTOR = {"DIR": "DIRECT", "MACS": "MACS", "GTS": "GTS", "BBC": "BBC", "7E": "7E", "FBC": "FBC", "ROMARA": "ROMARA (OLD LCI)"}
STAGE = {"Queue": "QUEUED", "Pricing": "PRICING", "Waiting": "AWAITING_SUPPLIER", "Submitted": "SUBMITTED"}
ESTIMATOR = {"RAMUS": "RAMUS", "GIORG": "GEORGIE"}
SKIP_REFS = {"ENQ-26-167", "ENQ-26-168"}          # Ramus's test rows
# rows that already exist in the app (from Tom's Excel) -> update instead of insert
MATCH = {("GHOBASH", "TWIN VILLA"): "26-009", ("DR. JOY", "JI C28 V13"): "26-095", ("MONTROSE", "FTR D58"): "2324"}

def q(v):
    if v is None or v == "": return "null"
    if isinstance(v, dt.datetime): return f"'{v.date()}'"
    return "'" + str(v).strip().replace("'", "''") + "'"

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb["Pipeline"]
hdr = [c.value for c in ws[1]]
rows = [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True) if any(v not in (None, "") for v in r)]

ins, upd, skipped, flags = [], [], [], []
for r in rows:
    ref = (r["Ref"] or "").strip()
    if ref in SKIP_REFS: skipped.append(f"{ref} {r['End User']} — test row"); continue
    enq = ref.replace("ENQ-", "") if ref else None
    end_user = (r["End User"] or "").strip() or None
    loc = (r["Location"] or "").strip() or None
    name = loc or end_user or "UNNAMED"
    stage = STAGE[str(r["Status"]).strip()]
    est = ESTIMATOR.get(str(r["Estimator"]).strip().upper()) if r["Estimator"] else None
    contractor = CONTRACTOR[str(r["Client"]).strip()]
    folder = f"Z:\\3_PROJECTS\\{r['Folder'].strip()}" if r["Folder"] else None
    waiting = (r["Waiting On"] or "").strip() or None
    if stage != "AWAITING_SUPPLIER" and waiting: flags.append(f"{name}: 'Waiting on {waiting}' but status {r['Status']} — kept the status, kept the note")
    key = (str(end_user).upper(), str(loc).upper())
    if key in MATCH:
        upd.append(f"update projects set quote_stage='{stage}', quote_type={q(r['Quote Type'])}, scope={q(r['Scope'])}, raised_by={q(r['Pending From'])}, "
                   f"quote_submitted_date=coalesce(quote_submitted_date,{q(r['Submitted'])}), estimator_id=coalesce(estimator_id,(select id from people where name={q(est)})) "
                   f"where enquiry_no='{MATCH[key]}';")
        continue
    if not ref: flags.append(f"{name} ({end_user or '—'}, {contractor}) has no ENQ number in Ramus's tracker")
    ins.append("(" + ", ".join([
        q(enq), q(name), q(end_user),
        f"(select id from contractors where name={q(contractor)})",
        "'QUOTE'", f"'{stage}'",
        f"(select id from people where name={q(est)})" if est else "null",
        q(r["Received"]), q(r["Due"]), q(r["Submitted"]),
        q(r["Quote Type"]), q(r["Scope"]), q(waiting), q(r["Pending From"]), q(folder), q(r["Notes"]),
    ]) + ")")

sql = ["set search_path = scheduler, public;", "begin;"]
sql += upd
sql.append("insert into projects (enquiry_no, name, client, contractor_id, status, quote_stage, estimator_id, received_date, quote_due_date, quote_submitted_date, quote_type, scope, waiting_on, raised_by, folder_path, notes) values")
sql.append(",\n".join(ins) + ";")
sql.append("commit;")
print("\n".join(sql))

by = collections.Counter(STAGE[str(r["Status"]).strip()] for r in rows if (r["Ref"] or "") not in SKIP_REFS)
rep = [f"# Pipeline import — {dt.date.today()}", "", f"Source: {SRC.split('/')[-1]}, sheet Pipeline, {len(rows)} rows.", "",
       f"- Inserted: {len(ins)}", f"- Updated (already in the app): {len(upd)}", f"- Skipped: {len(skipped)}", "",
       "## By stage", *[f"- {k}: {v}" for k, v in by.items()], "",
       "## Updated (matched to rows already in the app)", *[f"- {k[0]} / {k[1]} → {v}" for k, v in MATCH.items()], "",
       "## Skipped", *[f"- {s}" for s in skipped], "",
       "## Flags", *[f"- {f}" for f in flags]]
open("pipeline_report.md", "w").write("\n".join(rep))
